#!/usr/bin/env python3
import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


def build_sample_rows():
    return [
        {"service": "Acces", "score": 92, "owner": "IT", "status": "OK"},
        {"service": "Sauvegardes", "score": 88, "owner": "Infra", "status": "OK"},
        {"service": "Patching", "score": 73, "owner": "SecOps", "status": "A surveiller"},
        {"service": "Journalisation", "score": 81, "owner": "SOC", "status": "OK"},
        {"service": "Chiffrement", "score": 65, "owner": "SecOps", "status": "A prioriser"},
    ]


def style_header(cells):
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in cells:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def autosize_columns(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = max_length + 2


def add_kpis(ws, rows):
    ws["A1"] = "Rapport de durcissement"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A3"] = "Date"
    ws["B3"] = date.today().isoformat()
    ws["A4"] = "Nombre de controles"
    ws["B4"] = len(rows)
    ws["A5"] = "Score moyen"
    ws["B5"] = round(sum(item["score"] for item in rows) / len(rows), 1)

    for cell in ("A3", "A4", "A5"):
        ws[cell].font = Font(bold=True)


def add_table(ws, rows):
    ws.append(["Service", "Score", "Responsable", "Statut"])
    style_header(ws[1])
    for row in rows:
        ws.append([row["service"], row["score"], row["owner"], row["status"]])

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "0"


def add_chart(ws, rows_count):
    chart = BarChart()
    chart.title = "Scores par service"
    chart.y_axis.title = "Score"
    chart.x_axis.title = "Service"
    data = Reference(ws, min_col=2, min_row=1, max_row=rows_count + 1)
    categories = Reference(ws, min_col=1, min_row=2, max_row=rows_count + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, "F2")


def generate_report(output_path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthese"

    add_kpis(ws, rows)

    ws.append([])
    ws.append([])

    add_table(ws, rows)
    autosize_columns(ws)
    add_chart(ws, len(rows))

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Generer un rapport Excel de durcissement.")
    parser.add_argument(
        "-o",
        "--output",
        default="rapport_durcissement.xlsx",
        help="Chemin du fichier XLSX genere.",
    )
    parser.add_argument(
        "--with-sample",
        action="store_true",
        help="Inclure des donnees d'exemple.",
    )
    args = parser.parse_args()

    rows = build_sample_rows() if args.with_sample else []
    if not rows:
        rows = [
            {"service": "A completer", "score": 0, "owner": "N/A", "status": "N/A"}
        ]

    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    generate_report(output_path, rows)
    print(f"Rapport genere: {output_path}")


if __name__ == "__main__":
    main()
